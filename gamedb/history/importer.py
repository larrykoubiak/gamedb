"""Import recovered GameDB generations as lossless identity evidence."""

from __future__ import annotations

import csv
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import sqlite3
from typing import Iterable, Iterator, Mapping, Sequence

from openpyxl import load_workbook
from sqlalchemy import create_engine, insert, text
from sqlalchemy.engine import Connection, Engine

from . import schema


LEGACY_SNAPSHOT = "legacy-gamedb-2016"
PYTHON_SNAPSHOT = "python-gamedb-2016-2017"
BATCH_SIZE = 5_000


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while block := handle.read(1024 * 1024):
            digest.update(block)
    return digest.hexdigest()


def artifact_row(snapshot_key: str, artifact_key: str, role: str, path: Path) -> dict:
    stat = path.stat()
    return {
        "snapshot_key": snapshot_key,
        "artifact_key": artifact_key,
        "role": role,
        "source_path": str(path),
        "size_bytes": stat.st_size,
        "modified_at": datetime.fromtimestamp(stat.st_mtime, timezone.utc),
        "sha256": sha256_file(path),
    }


def match_review_rows(path: Path) -> Iterator[dict]:
    with path.open(encoding="utf-8-sig", errors="replace", newline="") as handle:
        for row_number, row in enumerate(csv.DictReader(handle, delimiter=";"), 2):
            original = int(row["matchScore"])
            reviewed = int(row["NewScore"])
            if reviewed == 0 and original != 0:
                decision = "rejected"
            elif reviewed == 100 and original != 100:
                decision = "accepted"
            else:
                decision = "unreviewed"
            yield {
                "snapshot_key": PYTHON_SNAPSHOT,
                "row_number": row_number,
                "system_name": row["System"],
                "software_name": row["softwareName"],
                "normalized_software_name": row["normalizedSoftwareName"] or None,
                "matched_external_name": row["normalizedReleaseName"] or None,
                "match_method": row["matchMethod"] or None,
                "original_score": original,
                "external_url": row["URL"] or None,
                "reviewed_score": reviewed,
                "review_decision": decision,
            }


def _json_value(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime,)):
        return value.isoformat()
    return str(value)


def workbook_rows(snapshot_key: str, artifact_key: str, path: Path) -> Iterator[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header_values = next(rows, None)
            if header_values is None:
                continue
            headers = [
                str(value).strip() if value not in (None, "") else f"column_{index + 1}"
                for index, value in enumerate(header_values)
            ]
            for row_number, values in enumerate(rows, 2):
                if not any(value is not None for value in values):
                    continue
                document = {
                    headers[index]: _json_value(value)
                    for index, value in enumerate(values)
                    if value is not None
                }
                yield {
                    "snapshot_key": snapshot_key,
                    "artifact_key": artifact_key,
                    "sheet_name": sheet.title,
                    "row_number": row_number,
                    "values": document,
                }
    finally:
        workbook.close()


class HistoricalImporter:
    """Replace only history-owned staging tables in one transaction."""

    def __init__(
        self,
        target_engine: Engine,
        old_database_url: str,
        python_gamedb: Path,
        legacy_archive: Path | None = None,
        patch_workbook: Path | None = None,
    ):
        self.target_engine = target_engine
        self.old_database_url = old_database_url
        self.python_gamedb = python_gamedb
        self.legacy_archive = legacy_archive
        self.patch_workbook = patch_workbook

    def execute(self) -> dict[str, int]:
        sqlite_path = self.python_gamedb / "sqlite" / "GameDB.db"
        review_path = self.python_gamedb / "output" / "matchResults.csv"
        for path in (sqlite_path, review_path):
            if not path.is_file():
                raise FileNotFoundError(f"historical source not found: {path}")
        if self.legacy_archive is not None and not self.legacy_archive.is_file():
            raise FileNotFoundError(f"legacy archive not found: {self.legacy_archive}")
        if self.patch_workbook is not None and not self.patch_workbook.is_file():
            raise FileNotFoundError(f"patch workbook not found: {self.patch_workbook}")

        old_engine = create_engine(self.old_database_url, future=True)
        sqlite_connection = sqlite3.connect(f"file:{sqlite_path}?mode=ro", uri=True)
        sqlite_connection.row_factory = sqlite3.Row
        counts: dict[str, int] = {}
        try:
            with self.target_engine.begin() as target:
                target.execute(text("CREATE SCHEMA IF NOT EXISTS staging"))
                schema.metadata.create_all(target)
                # Early development versions made both map endpoints required.
                # The recovered release map contains deliberately unresolved
                # right-hand endpoints, which are evidence too.
                target.execute(text(
                    "ALTER TABLE staging.history_python_map "
                    "ALTER COLUMN left_id DROP NOT NULL, "
                    "ALTER COLUMN right_id DROP NOT NULL"
                ))
                self._clear(target)
                self._snapshots(target, sqlite_path)
                with old_engine.connect() as old:
                    counts.update(self._legacy(target, old))
                counts.update(self._python_database(target, sqlite_connection))
                counts["match_reviews"] = self._insert(
                    target, schema.match_review, match_review_rows(review_path)
                )
                counts["artifacts"] = self._artifacts(target, sqlite_path, review_path)
                counts["workbook_rows"] = self._workbooks(target)
                self._indexes(target)
        finally:
            sqlite_connection.close()
            old_engine.dispose()
        return counts

    @staticmethod
    def _clear(connection: Connection) -> None:
        names = ", ".join(
            f"staging.{table.name}"
            for table in reversed(schema.metadata.sorted_tables)
        )
        connection.execute(text(f"TRUNCATE {names}"))

    def _snapshots(self, connection: Connection, sqlite_path: Path) -> None:
        now = datetime.now(timezone.utc)
        archive_sha = sha256_file(self.legacy_archive) if self.legacy_archive else None
        connection.execute(insert(schema.snapshot), [
            {
                "snapshot_key": LEGACY_SNAPSHOT,
                "source_kind": "legacy relational database",
                "label": "GameDB hierarchy and sourced claims recovered from the 2016 database",
                "source_locator": "OLD_GAMEDB",
                "source_sha256": archive_sha,
                "details": {
                    "archive": str(self.legacy_archive) if self.legacy_archive else None,
                    "note": "The queryable copy matches the 2016 dump's core identity counts; claims may include later cleanup.",
                },
                "imported_at": now,
            },
            {
                "snapshot_key": PYTHON_SNAPSHOT,
                "source_kind": "SQLite working database and review artifacts",
                "label": "Parent/Clone import, external mappings, and reviewed fuzzy matches",
                "source_locator": str(self.python_gamedb),
                "source_sha256": sha256_file(sqlite_path),
                "details": {
                    "database": str(sqlite_path),
                    "note": "Generated grouping is preserved separately from explicit match review decisions.",
                },
                "imported_at": now,
            },
        ])

    def _legacy(self, target: Connection, old: Connection) -> dict[str, int]:
        queries = {
            schema.legacy_source: '''
                SELECT :snapshot_key AS snapshot_key,
                    "sourceId" AS legacy_source_id, "sourceName" AS source_name,
                    "sourceURL" AS source_url, "sourceOrder" AS source_order,
                    CAST("hasSoftwareData" AS integer) AS has_software_data,
                    CAST("hasDATinfo" AS integer) AS has_dat_info,
                    CAST("hasReleaseData" AS integer) AS has_release_data,
                    CAST("hasImages" AS integer) AS has_images
                FROM tblsources ORDER BY "sourceId"
            ''',
            schema.legacy_system: '''
                SELECT :snapshot_key AS snapshot_key,
                    "systemId" AS legacy_system_id, "systemName" AS system_name,
                    "systemManufacturer" AS manufacturer_name,
                    "systemFormat" AS system_format, "systemAcronym" AS system_acronym
                FROM tblsystems ORDER BY "systemId"
            ''',
            schema.legacy_software: '''
                SELECT :snapshot_key AS snapshot_key,
                    s."softwareId" AS legacy_software_id,
                    s."softwareName" AS software_name, s."softwareType" AS software_type,
                    s."softwareTypeId" AS software_type_id,
                    s."systemId" AS legacy_system_id, st."statusName" AS status_name
                FROM tblsoftwares s LEFT JOIN tblstatus st ON st."statusId"=s."statusId"
                ORDER BY s."softwareId"
            ''',
            schema.legacy_release: '''
                SELECT :snapshot_key AS snapshot_key,
                    r."releaseId" AS legacy_release_id, r."releaseName" AS release_name,
                    r."releaseRegion" AS region, r."releaseType" AS release_type,
                    r."softwareId" AS legacy_software_id, st."statusName" AS status_name,
                    r."releaseFolder" AS release_folder
                FROM tblreleases r LEFT JOIN tblstatus st ON st."statusId"=r."statusId"
                ORDER BY r."releaseId"
            ''',
            schema.legacy_file: '''
                SELECT :snapshot_key AS snapshot_key,
                    "fileId" AS legacy_file_id, "releaseId" AS legacy_release_id,
                    "fileName" AS file_name, "fileSize" AS size,
                    lower("fileCRC") AS crc_hex, lower("fileMD5") AS md5_hex,
                    lower("fileSHA1") AS sha1_hex, "fileStatus" AS file_status
                FROM tblfiles ORDER BY "fileId"
            ''',
        }
        counts = {}
        for table, statement in queries.items():
            rows = old.execute(text(statement), {"snapshot_key": LEGACY_SNAPSHOT}).mappings()
            counts[table.name] = self._insert(target, table, rows)

        claim_statement = text('''
            SELECT :snapshot_key AS snapshot_key, 'software' AS subject_type,
                d."softwareDataId" AS legacy_claim_id, d."softwareId" AS subject_id,
                f."softwareFlagName" AS flag_name, d."softwareFlagValue" AS flag_value,
                d."softwareFlagSource" AS source_name, d."sourceId" AS source_record_id
            FROM tblsoftwaredata d
            JOIN tblsoftwareflags f ON f."softwareFlagId"=d."softwareFlagId"
            UNION ALL
            SELECT :snapshot_key, 'release', d."releaseDataId", d."releaseId",
                f."releaseFlagName", d."releaseFlagValue",
                d."releaseFlagSource", d."sourceId"
            FROM tblreleasedata d
            JOIN tblreleaseflags f ON f."releaseFlagId"=d."releaseFlagId"
            ORDER BY subject_type, legacy_claim_id
        ''')
        counts[schema.legacy_claim.name] = self._insert(
            target, schema.legacy_claim,
            old.execute(claim_statement, {"snapshot_key": LEGACY_SNAPSHOT}).mappings(),
        )
        rating_statement = text('''
            SELECT :snapshot_key AS snapshot_key, 'software' AS subject_type,
                r."sourceSoftwareFlagRatingId" AS legacy_rating_id,
                r."sourceId" AS legacy_source_id,
                r."softwareFlagId" AS legacy_flag_id,
                f."softwareFlagName" AS flag_name, r.rating
            FROM tblsourcesoftwareflagrating r
            LEFT JOIN tblsoftwareflags f
              ON f."softwareFlagId"=r."softwareFlagId"
            UNION ALL
            SELECT :snapshot_key, 'release',
                r."sourceReleaseFlagRatingId", r."sourceId",
                r."releaseFlagId", f."releaseFlagName", r.rating
            FROM tblsourcereleaseflagrating r
            LEFT JOIN tblreleaseflags f
              ON f."releaseFlagId"=r."releaseFlagId"
            ORDER BY subject_type, legacy_rating_id
        ''')
        counts[schema.legacy_source_rating.name] = self._insert(
            target, schema.legacy_source_rating,
            old.execute(
                rating_statement, {"snapshot_key": LEGACY_SNAPSHOT}
            ).mappings(),
        )
        return counts

    def _python_database(
        self, target: Connection, source: sqlite3.Connection
    ) -> dict[str, int]:
        specs = (
            (schema.python_dat_file, '''SELECT ? snapshot_key, datFileId dat_file_id,
                systemId legacy_system_id, datFileName file_name, datType dat_type,
                datReleaseGroup release_group, datDate dat_date FROM tblDatFiles ORDER BY datFileId'''),
            (schema.python_dat_game, '''SELECT ? snapshot_key, datGameId dat_game_id,
                datFileId dat_file_id, datGameName game_name, datCloneOf clone_of,
                datRomOf rom_of FROM tblDatGames ORDER BY datGameId'''),
            (schema.python_dat_rom, '''SELECT ? snapshot_key, datROMId dat_rom_id,
                datFileId dat_file_id, datGameId dat_game_id, datROMName rom_name,
                datROMMerge merge_name, datROMSize size, lower(datROMCRC) crc_hex,
                lower(datROMMD5) md5_hex, lower(datROMSHA1) sha1_hex
                FROM tblDatRoms ORDER BY datROMId'''),
            (schema.python_system, '''SELECT ? snapshot_key, systemId legacy_system_id,
                systemName system_name, systemManufacturer manufacturer_name
                FROM tblSystems ORDER BY systemId'''),
            (schema.python_software, '''SELECT ? snapshot_key, softwareId legacy_software_id,
                softwareName software_name, softwareType software_type,
                systemId legacy_system_id FROM tblSoftwares ORDER BY softwareId'''),
            (schema.python_release, '''SELECT ? snapshot_key, releaseId legacy_release_id,
                releaseName release_name, releaseType release_type,
                softwareId legacy_software_id FROM tblReleases ORDER BY releaseId'''),
            (schema.python_rom, '''SELECT ? snapshot_key, romId legacy_rom_id,
                releaseId legacy_release_id, size, lower(crc32) crc_hex,
                lower(md5) md5_hex, lower(sha1) sha1_hex FROM tblROMs ORDER BY romId'''),
            (schema.python_external_game, '''SELECT ? snapshot_key,
                g.scraperGameId external_game_id, s.scraperName source_name,
                ss.scraperSystemId external_system_id,
                ss.scraperSystemName external_system_name,
                ss.scraperSystemURL external_system_url,
                g.scraperGameName game_name, g.scraperGameURL game_url
                FROM tblScraperGames g
                JOIN tblScraperSystems ss ON ss.scraperSystemId=g.scraperSystemId
                JOIN tblScrapers s ON s.scraperId=ss.scraperId
                ORDER BY g.scraperGameId'''),
            (schema.python_external_release, '''SELECT ? snapshot_key,
                scraperReleaseId external_release_id, scraperGameId external_game_id,
                scraperReleaseName release_name, scraperReleaseRegion region,
                scraperReleaseType release_type FROM tblScraperReleases
                ORDER BY scraperReleaseId'''),
        )
        counts = {}
        for table, query in specs:
            cursor = source.execute(query, (PYTHON_SNAPSHOT,))
            rows = (dict(row) for row in cursor)
            counts[table.name] = self._insert(target, table, rows)

        def map_rows(map_type: str, query: str) -> Iterator[dict]:
            for row in source.execute(query):
                yield {
                    "snapshot_key": PYTHON_SNAPSHOT,
                    "map_type": map_type,
                    "left_id": row[0],
                    "right_id": row[1],
                }

        maps = (
            ("system_external", "SELECT systemId, scraperSystemId FROM tblSystemMap"),
            ("software_external", "SELECT softwareId, scraperGameId FROM tblSoftwareMap"),
            ("release_external", "SELECT releaseId, scraperReleaseId FROM tblReleaseMap"),
        )
        mapping_rows = (
            row for map_type, query in maps for row in map_rows(map_type, query)
        )
        counts[schema.python_map.name] = self._insert(
            target, schema.python_map,
            (
                {**row, "row_number": row_number}
                for row_number, row in enumerate(mapping_rows, 1)
            ),
        )
        counts[schema.python_synonym.name] = self._insert(
            target, schema.python_synonym,
            ({
                "snapshot_key": PYTHON_SNAPSHOT,
                "synonym_type": row[2] or "",
                "source_value": row[0] or "",
                "normalized_value": row[1] or "",
            } for row in source.execute("SELECT key, value, type FROM tblSynonyms")),
        )
        return counts

    def _artifacts(
        self, connection: Connection, sqlite_path: Path, review_path: Path
    ) -> int:
        rows = [
            artifact_row(PYTHON_SNAPSHOT, "python-sqlite", "historical derived database", sqlite_path),
            artifact_row(PYTHON_SNAPSHOT, "match-review", "manual match review", review_path),
        ]
        gamefaqs = self.python_gamedb / "sqlite" / "GameFAQs.db"
        if gamefaqs.is_file():
            rows.append(artifact_row(
                PYTHON_SNAPSHOT, "gamefaqs-sqlite", "external source cache", gamefaqs
            ))
        for dat_path in sorted((self.python_gamedb / "DAT").glob("*.dat")):
            rows.append(artifact_row(
                PYTHON_SNAPSHOT, f"dat:{dat_path.name}", "Parent/Clone DAT", dat_path
            ))
        if self.legacy_archive:
            rows.append(artifact_row(
                LEGACY_SNAPSHOT, "mysql-backup", "immutable SQL backup archive",
                self.legacy_archive,
            ))
        if self.patch_workbook:
            rows.append(artifact_row(
                PYTHON_SNAPSHOT, "later-patches", "manual mapping patches",
                self.patch_workbook,
            ))
        genre = self.python_gamedb / "GameFAQs Genre WIP - 14-11-2016.xlsx"
        if genre.is_file():
            rows.append(artifact_row(
                PYTHON_SNAPSHOT, "genre-review", "manual metadata vocabulary review", genre
            ))
        snes = self.python_gamedb / "output" / "snes_matchResult.xlsx"
        if snes.is_file():
            rows.append(artifact_row(
                PYTHON_SNAPSHOT, "snes-match-review", "historical match analysis", snes
            ))
        return self._insert(connection, schema.artifact, rows)

    def _workbooks(self, connection: Connection) -> int:
        inputs: list[tuple[str, Path]] = []
        if self.patch_workbook:
            inputs.append(("later-patches", self.patch_workbook))
        genre = self.python_gamedb / "GameFAQs Genre WIP - 14-11-2016.xlsx"
        if genre.is_file():
            inputs.append(("genre-review", genre))
        snes = self.python_gamedb / "output" / "snes_matchResult.xlsx"
        if snes.is_file():
            inputs.append(("snes-match-review", snes))
        return self._insert(
            connection, schema.workbook_row,
            (row for key, path in inputs for row in workbook_rows(PYTHON_SNAPSHOT, key, path)),
        )

    @staticmethod
    def _insert(connection: Connection, table, rows: Iterable[Mapping]) -> int:
        count = 0
        batch = []
        for row in rows:
            batch.append(dict(row))
            if len(batch) >= BATCH_SIZE:
                connection.execute(insert(table), batch)
                count += len(batch)
                batch.clear()
        if batch:
            connection.execute(insert(table), batch)
            count += len(batch)
        return count

    @staticmethod
    def _indexes(connection: Connection) -> None:
        statements = (
            "CREATE INDEX IF NOT EXISTS history_legacy_file_sha1_idx ON staging.history_legacy_file (sha1_hex)",
            "CREATE INDEX IF NOT EXISTS history_python_rom_sha1_idx ON staging.history_python_rom (sha1_hex)",
            "CREATE INDEX IF NOT EXISTS history_python_dat_rom_sha1_idx ON staging.history_python_dat_rom (sha1_hex)",
            "CREATE INDEX IF NOT EXISTS history_legacy_release_software_idx ON staging.history_legacy_release (legacy_software_id)",
            "CREATE INDEX IF NOT EXISTS history_python_release_software_idx ON staging.history_python_release (legacy_software_id)",
            "CREATE INDEX IF NOT EXISTS history_python_map_right_idx ON staging.history_python_map (map_type, right_id)",
            "CREATE INDEX IF NOT EXISTS history_match_review_decision_idx ON staging.history_match_review (review_decision)",
        )
        for statement in statements:
            connection.execute(text(statement))
